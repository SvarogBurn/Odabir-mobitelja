# -*- coding: utf-8 -*-
"""
Sastavlja jedinstveni seminarski .docx iz markdown poglavlja, s ugradjenim
grafovima. Redoslijed: Uvod + Analiza (1-2) -> Interpretacija (3) -> Zakljucak +
Literatura (4-5). Grafovi se umecu kod naslova koji ih spominju.

Pokretanje:  python build_docx.py
Rezultat:    Seminar_AHP_Odabir_mobitela.docx
"""

import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "Seminar_AHP_Odabir_mobitela.docx"
IMGDIR = "output"

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+`)')
PNG = re.compile(r'`([^`]+\.png)`')

# Opisni naslovi slika (numeriranje se dodjeljuje automatski redom umetanja)
FIG_CAPTIONS = {
    '0_hijerarhija.png': 'Hijerarhijski model problema odlučivanja '
        '(cilj → kriteriji → podkriteriji → alternative)',
    '1_performance.png': 'Performance — lokalni prioriteti alternativa po '
        'kriterijima u odnosu na težine kriterija',
    '2_dynamic_base.png': 'Dynamic — trenutne težine kriterija (lijevo) i '
        'ukupni prioriteti alternativa (desno)',
    '2_dynamic_scenarios.png': 'Dynamic — utjecaj promjene težine kriterija za '
        '±10 % na ukupne prioritete alternativa',
    '2_dynamic_components.png': 'Dynamic / Components — doprinos pojedinih '
        'kriterija ukupnom prioritetu svake alternative',
    '3_gradient.png': 'Gradient — promjena prioriteta alternativa s promjenom '
        'težine pojedinog kriterija',
    '4_head_to_head.png': 'Head-to-head — usporedba dviju najboljih alternativa '
        'po svim listnim kriterijima',
    '5_two_d.png': '2D — položaj alternativa s obzirom na dva najvažnija '
        'kriterija (kvaliteta kamere i performanse)',
}


def add_runs(par, text, bold=False):
    """Dodaje tekst u paragraf uz **bold** i `monospace` formatiranje."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()]); r.bold = bold
        tok = m.group(0)
        if tok.startswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('*'):
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:]); r.bold = bold


class Builder:
    def __init__(self, doc):
        self.doc = doc
        self.title_used = False
        self.fig_no = 0        # brojac slika
        self.buf = []          # tekuci blok (paragraf/stavka liste)
        self.buf_kind = None   # 'p' | 'bullet' | 'number'

    # ---- upravljanje tekucim blokom ----
    def flush(self):
        if not self.buf:
            return
        text = " ".join(self.buf).strip()
        self.buf = []
        kind, self.buf_kind = self.buf_kind, None
        if not text:
            return
        if kind == 'bullet':
            p = self.doc.add_paragraph(style='List Bullet')
        elif kind == 'number':
            p = self.doc.add_paragraph(style='List Number')
        else:
            p = self.doc.add_paragraph()
        add_runs(p, text)

    def append(self, text, kind):
        if self.buf and self.buf_kind == kind and kind == 'p':
            self.buf.append(text)
        else:
            self.flush()
            self.buf = [text]
            self.buf_kind = kind

    def cont(self, text):
        """Nastavak (omotani redak) tekuceg bloka."""
        if self.buf:
            self.buf.append(text)
        else:
            self.append(text, 'p')

    # ---- elementi ----
    def heading(self, level, text):
        self.flush()
        pngs = PNG.findall(text)
        text = re.sub(r'\s*[—,]?\s*`[^`]+\.png`', '', text).strip()
        if level == 1 and not self.title_used:
            self.title_used = True
            self.doc.add_heading(text, level=0)
        else:
            self.doc.add_heading(text, level=level)
        if len(pngs) == 1:
            self.image(pngs[0])

    def image(self, fname):
        path = os.path.join(IMGDIR, fname)
        if not os.path.exists(path):
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Inches(6.2))
        self.fig_no += 1
        title = FIG_CAPTIONS.get(fname, fname)
        try:
            cap = self.doc.add_paragraph(style='Caption')
        except KeyError:
            cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Slika {self.fig_no}. {title}")
        r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    def code(self, lines):
        self.flush()
        p = self.doc.add_paragraph()
        for i, ln in enumerate(lines):
            if i:
                p.add_run().add_break()
            r = p.add_run(ln); r.font.name = 'Consolas'; r.font.size = Pt(8.5)

    def caption(self, text):
        self.flush()
        try:
            p = self.doc.add_paragraph(style='Caption')
        except KeyError:
            p = self.doc.add_paragraph()
        add_runs(p, text)

    def table(self, rows):
        self.flush()
        cells = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
        header, body = cells[0], cells[2:]
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = 'Light Grid Accent 1'
        for j, h in enumerate(header):
            add_runs(t.rows[0].cells[j].paragraphs[0], h, bold=True)
        for row in body:
            tc = t.add_row().cells
            for j, val in enumerate(row[:len(header)]):
                add_runs(tc[j].paragraphs[0], val)
        self.doc.add_paragraph()


def title_page(doc):
    """Naslovna stranica seminara."""
    def center(text, size=11, bold=False, italic=False, after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.bold, r.italic, r.font.size = bold, italic, Pt(size)
        return p

    center("Sveučilište u Rijeci", 13, bold=True, after=0)
    center("Fakultet informatike i digitalnih tehnologija", 12, after=0)
    center("Sveučilišni diplomski studij Informatika", 11, italic=True)
    for _ in range(4):
        doc.add_paragraph()
    center("Reo Turčinović", 14, bold=True)
    for _ in range(3):
        doc.add_paragraph()
    center("ODABIR OPTIMALNOG PAMETNOG TELEFONA", 22, bold=True, after=0)
    center("PRIMJENOM AHP METODE", 22, bold=True, after=8)
    center("Seminarski rad", 12, italic=True, after=0)
    center("Kolegij: Upravljanje digitalnom transformacijom", 11)
    for _ in range(6):
        doc.add_paragraph()
    center("Mentori:", 11, after=0)
    center("prof. dr. sc. Patrizia Poščić", 11, after=0)
    center("doc. dr. sc. Kristian Stančin", 11)
    for _ in range(4):
        doc.add_paragraph()
    center("Rijeka, lipanj 2026.", 11)
    doc.add_page_break()


def parse(builder, text):
    lines = text.split('\n')
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        # code fence
        if s.startswith('```'):
            block = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                block.append(lines[i]); i += 1
            builder.code(block); i += 1; continue
        # table
        if s.startswith('|'):
            block = []
            while i < n and lines[i].strip().startswith('|'):
                block.append(lines[i]); i += 1
            builder.table(block); continue
        # blank
        if s == '':
            builder.flush(); i += 1; continue
        # horizontal rule
        if s == '---':
            builder.flush(); i += 1; continue
        # table caption (Tablica N. ...)
        if re.match(r'^Tablica \d+\.', s):
            builder.caption(s); i += 1; continue
        # blockquote (struktura-napomena) -> preskoci
        if s.startswith('>'):
            builder.flush(); i += 1; continue
        # headings
        if s.startswith('### '):
            builder.heading(3, s[4:]); i += 1; continue
        if s.startswith('## '):
            builder.heading(2, s[3:]); i += 1; continue
        if s.startswith('# '):
            builder.heading(1, s[2:]); i += 1; continue
        # bullet
        if s.startswith('- '):
            builder.append(s[2:], 'bullet'); i += 1; continue
        # numbered
        m = re.match(r'^\d+\.\s+(.*)', s)
        if m:
            builder.append(m.group(1), 'number'); i += 1; continue
        # standalone image ![](path)
        mi = re.match(r'!\[.*?\]\((.+?)\)', s)
        if mi:
            builder.flush(); builder.image(os.path.basename(mi.group(1))); i += 1; continue
        # plain / continuation
        if line.startswith((' ', '\t')) and builder.buf:
            builder.cont(s)
        else:
            builder.append(s, 'p')
        i += 1
    builder.flush()


def main():
    seminar = open("Seminar_Uvod_Analiza_Zakljucak.md", encoding='utf-8').read()
    interp = open("Interpretacija_rezultata.md", encoding='utf-8').read()

    # podijeli seminar na (Uvod+Analiza) i (Zakljucak+Literatura)
    idx = seminar.index("# 4. Zaključak")
    head, tail = seminar[:idx], seminar[idx:]
    # ukloni markdown naslov (naslovna stranica ga vec nosi)
    head = re.sub(r'^# .*\n', '', head, count=1)

    doc = Document()
    # osnovni font
    style = doc.styles['Normal']
    style.font.name = 'Calibri'; style.font.size = Pt(11)

    title_page(doc)       # naslovna stranica

    b = Builder(doc)
    b.title_used = True   # nema vise markdown naslova; sve sekcije su Heading 1
    parse(b, head)        # 1. Uvod + 2. Analiza
    # podnaslov sa studentima nakon naslova nije moguc retroaktivno; preskoceno
    parse(b, interp)      # 3. Interpretacija (s grafovima)
    parse(b, tail)        # 4. Zakljucak + 5. Literatura

    doc.save(OUT)
    print(f"Spremljeno: {OUT}")


if __name__ == "__main__":
    main()
